"""
Turns a report from workflows.py into files: PDF (fpdf2), Word (python-docx),
Markdown and LaTeX (zips with their charts and references.bib), BibTeX, and the
data zip (the same as for chat answers: full series + what was fetched).
Every format is built from the same report dict, so they always agree.
"""

import io
import re
import zipfile
from pathlib import Path

IMAGE_TYPES = {".png", ".gif"}          # what PDF and Word embed (SVG fallback charts: web/Markdown only)
PDF_FONTS = [  # (regular, bold, italic): Unicode TrueType fonts, first found wins
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", None, None),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
]


# ------------------------------------------------------------------ text helpers
def blocks(markdown: str) -> list[tuple[str, str]]:
    """The summary's simple markdown as ("p" | "li", text) blocks."""
    out = []
    for para in re.split(r"\n\s*\n", markdown.strip()):
        lines = [l.strip() for l in para.splitlines() if l.strip()]
        if lines and all(re.match(r"([-*+]|\d+[.)])\s", l) for l in lines):
            out += [("li", re.sub(r"^([-*+]|\d+[.)])\s+", "", l)) for l in lines]
        elif lines:
            out.append(("p", " ".join(lines)))
    return out


def runs(text: str) -> list[tuple[str, bool]]:
    """(text, bold) pieces of a line with **bold** marks (other marks dropped)."""
    text = re.sub(r"(?<![*\w])\*([^*\s][^*]*)\*(?!\*)", r"\1", text).replace("`", "")
    return [(part, i % 2 == 1) for i, part in enumerate(text.split("**")) if part]


def images(section: dict) -> list[Path]:
    return [Path(p) for p in section.get("charts", []) if Path(p).suffix in IMAGE_TYPES and Path(p).exists()]


def align(table: dict) -> list[str]:
    """Column alignment, "l" or "r": a table may set its own (text columns); default numbers right."""
    return table.get("align") or ["l"] + ["r"] * (len(table["columns"]) - 1)


def source_text(s: dict) -> str:
    """One line per source, the same in every format."""
    parts = [s["author"].strip("{}").replace(" and ", ", ") if s.get("author") else "",
             s.get("title", ""), s.get("journal", ""), s.get("year", "") if not s.get("accessed") else "",
             s.get("url", ""), f"accessed {s['accessed']}" if s.get("accessed") else "", s.get("note", "")]
    return ". ".join(p for p in parts if p) + "."


def summary_status(report: dict) -> str:
    s = report["summary"]
    if s["status"] == "ok":
        return "Every number in this summary was found in the data below (checked automatically; wording is not checked)."
    if s["status"] == "unverified":
        return f"Not found in the data (treat with caution): {', '.join(s['unverified'])}."
    if s["status"] == "python":
        return ("Summary listed by Python from the tables (no model)"
                + (f"; the model was unavailable: {s['reason']}" if s.get("reason") else "") + ".")
    return f"No written summary: {s.get('reason', 'the model was unavailable')}."


# ------------------------------------------------------------------ BibTeX
def bibtex(report: dict) -> str:
    entries = []
    for s in report["sources"]:
        fields = {k: s[k] for k in ("author", "title", "year", "journal", "url", "note") if s.get(k)}
        if s.get("accessed"):
            fields["urldate"] = s["accessed"]
            fields.setdefault("year", s["accessed"][:4])
        body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items())
        entries.append(f"@{s.get('type', 'misc')}{{{s['key']},\n{body}\n}}")
    return "\n\n".join(entries) + "\n"


# ------------------------------------------------------------------ Markdown
def markdown(report: dict, image_dir: str = "charts") -> str:
    out = [f"# {report['title']}", "", f"*Made {report['created'].replace('T', ' ')}*", ""]
    if report["summary"]["text"]:
        out += ["## Summary", "", report["summary"]["text"], ""]
    out += [f"> {summary_status(report)}", ""]
    for sec in report["sections"]:
        out += [f"## {sec['heading']}", ""]
        if sec.get("table"):
            cols = sec["table"]["columns"]
            out += ["| " + " | ".join(cols) + " |", "|" + "|".join("---:" if a == "r" else "---"
                                                                    for a in align(sec["table"])) + "|"]
            out += ["| " + " | ".join(str(c).replace("|", "/") for c in row) + " |" for row in sec["table"]["rows"]]
            out.append("")
        out += [f"- **[{i['title']}]({i['url']})** — {i['text']}" if i.get("url") else f"- **{i['title']}** — {i['text']}"
                for i in sec.get("items", [])] + ([""] if sec.get("items") else [])
        out += [f"- {n}" for n in sec["notes"]] + ([""] if sec["notes"] else [])
        out += [f"![{Path(p).stem}]({image_dir}/{Path(p).name})" for p in sec.get("charts", [])]
        out.append("")
    out += ["## About this report", ""] + [f"- {a}" for a in report["about"]] + [""]
    out += ["## Sources", ""] + [f"- {source_text(s)}" for s in report["sources"]]
    return "\n".join(out) + "\n"


# ------------------------------------------------------------------ LaTeX
_TEX = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{",
        "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}", "≥": r"$\geq$", "≤": r"$\leq$",
        "–": "--", "—": "---"}


def tex(text) -> str:
    return "".join(_TEX.get(ch, ch) for ch in str(text))


def tex_runs(text: str) -> str:
    return "".join(rf"\textbf{{{tex(t)}}}" if bold else tex(t) for t, bold in runs(text))


def latex(report: dict, image_dir: str = "figures") -> str:
    out = [r"\documentclass[11pt]{article}", r"% Compile with pdflatex, xelatex or lualatex, then bibtex.",
           r"\usepackage[utf8]{inputenc}", r"\usepackage[T1]{fontenc}", r"\usepackage[margin=2.5cm]{geometry}",
           r"\usepackage{booktabs,graphicx,hyperref}", "",
           rf"\title{{{tex(report['title'])}}}", rf"\date{{{tex(report['created'][:10])}}}", r"\author{}",
           r"\begin{document}", r"\maketitle", ""]
    if report["summary"]["text"]:
        out.append(r"\section*{Summary}")
        items = False
        for kind, text in blocks(report["summary"]["text"]):
            if kind == "li" and not items:
                out.append(r"\begin{itemize}")
            if kind != "li" and items:
                out.append(r"\end{itemize}")
            items = kind == "li"
            out.append((r"\item " if kind == "li" else "") + tex_runs(text) + ("" if kind == "li" else "\n"))
        if items:
            out.append(r"\end{itemize}")
    out += [rf"\noindent\emph{{{tex(summary_status(report))}}}", ""]
    for sec in report["sections"]:
        out.append(rf"\section*{{{tex(sec['heading'])}}}")
        if sec.get("table"):
            cols = sec["table"]["columns"]
            out += [r"\begin{table}[h]\centering\small",
                    r"\resizebox{\textwidth}{!}{%",
                    r"\begin{tabular}{p{6cm}" + "".join(a if a == "r" else "p{4cm}" for a in align(sec["table"])[1:]) + "}",
                    r"\toprule", " & ".join(tex(c) for c in cols) + r" \\", r"\midrule"]
            out += [" & ".join(tex(c) for c in row) + r" \\" for row in sec["table"]["rows"]]
            out += [r"\bottomrule", r"\end{tabular}}", r"\end{table}"]
        if sec.get("items"):
            out += [r"\begin{itemize}"] + [
                rf"\item \textbf{{{tex(i['title'])}}} --- {tex(i['text'])}"
                + (rf" \url{{{i['url']}}}" if i.get("url") else "") for i in sec["items"]] + [r"\end{itemize}"]
        if sec["notes"]:
            out += [r"\begin{itemize}\small"] + [rf"\item {tex(n)}" for n in sec["notes"]] + [r"\end{itemize}"]
        for p in images(sec):
            if p.suffix == ".png":
                out.append(rf"\begin{{center}}\includegraphics[width=0.85\textwidth]{{{image_dir}/{p.name}}}\end{{center}}")
        out.append("")
    out += [r"\section*{About this report}", r"\begin{itemize}"] + [rf"\item {tex(a)}" for a in report["about"]]
    out += [r"\end{itemize}", "Data: " + ", ".join(rf"\cite{{{s['key']}}}" for s in report["sources"]) + ".",
            r"\bibliographystyle{plain}", r"\bibliography{references}", r"\end{document}"]
    return "\n".join(out) + "\n"


# ------------------------------------------------------------------ PDF
def _pdf_font(pdf) -> str:
    for regular, bold, italic in PDF_FONTS:
        if Path(regular).exists():
            pdf.add_font("body", "", regular)
            pdf.add_font("body", "B", bold if bold and Path(bold).exists() else regular)
            pdf.add_font("body", "I", italic if italic and Path(italic).exists() else regular)
            return "body"
    return "helvetica"  # no Unicode font found: characters outside Latin-1 become "?"


def pdf(report: dict) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    doc = FPDF(format="A4")
    doc.set_auto_page_break(True, margin=18)
    doc.set_margins(18, 18, 18)
    font = _pdf_font(doc)
    dashes = str.maketrans({"\u2010": "-", "\u2011": "-", "\u2012": "-"})  # hyphens Arial lacks
    safe = (lambda s: str(s).translate(dashes)) if font == "body" else \
        (lambda s: str(s).translate(dashes).encode("latin-1", "replace").decode("latin-1"))
    doc.add_page()
    width = doc.epw

    def para(text: str, size=10.5, style="", gap=2.5, bullet=False):
        doc.set_font(font, style, size)
        doc.set_x(doc.l_margin)
        if bullet:
            doc.cell(5, 5.5, "•")
        line = "".join(f"**{t}**" if b else t.replace("*", "") for t, b in runs(text))
        doc.multi_cell(width - (5 if bullet else 0), 5.5 if size < 12 else 7, safe(line), markdown=True,
                       align="L", new_x="LMARGIN", new_y="NEXT")
        doc.ln(gap)

    para(report["title"], 18, "B", 1)
    para(f"Made {report['created'].replace('T', ' ')}", 9, "I", 4)
    if report["summary"]["text"]:
        para("Summary", 13, "B", 1)
        for kind, text in blocks(report["summary"]["text"]):
            para(text, bullet=kind == "li", gap=1 if kind == "li" else 2.5)
    para(summary_status(report), 9, "I", 4)
    for sec in report["sections"]:
        para(sec["heading"], 13, "B", 1)
        if sec.get("table"):
            cols = sec["table"]["columns"]
            first = min(70, width * 0.4)
            doc.set_font(font, "", 8.5)
            with doc.table(col_widths=[first] + [(width - first) / (len(cols) - 1)] * (len(cols) - 1),
                           text_align=["RIGHT" if a == "r" else "LEFT" for a in align(sec["table"])], line_height=4.5,
                           headings_style=FontFace(emphasis="BOLD"), borders_layout="HORIZONTAL_LINES") as table:
                for row in [cols] + sec["table"]["rows"]:
                    r = table.row()
                    for c in row:
                        r.cell(safe(c))
            doc.ln(2)
        for i in sec.get("items", []):
            para(f"**{i['title']}** — {i['text']}" + (f" {i['url']}" if i.get("url") else ""), 9, "", 1.5, bullet=True)
        for n in sec["notes"]:
            para(n, 8.5, "", 1, bullet=True)
        for p in images(sec):
            from PIL import Image
            with Image.open(p) as im:
                h = width * 0.8 * im.height / im.width
            if doc.get_y() + h > doc.page_break_trigger:
                doc.add_page()
            doc.image(str(p), x=doc.l_margin + width * 0.1, w=width * 0.8)
            doc.ln(3)
        doc.ln(2)
    para("About this report", 13, "B", 1)
    for a in report["about"]:
        para(a, 9, "", 1, bullet=True)
    para("Sources", 13, "B", 1)
    for s in report["sources"]:
        para(source_text(s), 9, "", 1)
    return bytes(doc.output())


# ------------------------------------------------------------------ Word
def docx(report: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(report["title"], 0)
    doc.add_paragraph().add_run(f"Made {report['created'].replace('T', ' ')}").italic = True

    def rich(paragraph, text):
        for t, bold in runs(text):
            paragraph.add_run(t).bold = bold

    if report["summary"]["text"]:
        doc.add_heading("Summary", 1)
        for kind, text in blocks(report["summary"]["text"]):
            rich(doc.add_paragraph(style="List Bullet" if kind == "li" else None), text)
    doc.add_paragraph().add_run(summary_status(report)).italic = True
    for sec in report["sections"]:
        doc.add_heading(sec["heading"], 1)
        for i in sec.get("items", []):
            item = doc.add_paragraph(style="List Bullet")
            item.add_run(i["title"]).bold = True
            item.add_run(f" — {i['text']}" + (f" {i['url']}" if i.get("url") else ""))
        if not sec.get("table"):
            for n in sec["notes"]:
                doc.add_paragraph(n, style="List Bullet").runs[0].font.size = Pt(9)
            continue
        cols, rows = sec["table"]["columns"], sec["table"]["rows"]
        table = doc.add_table(rows=1 + len(rows), cols=len(cols))
        table.style = "Light List Accent 1"
        first = Inches(2.4)
        rest = Inches((6.3 - 2.4) / max(1, len(cols) - 1))
        for i, row in enumerate([cols] + rows):
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.width = first if j == 0 else rest  # explicit widths: some viewers ignore table styles
                cell.text = str(value)
                para = cell.paragraphs[0]
                para.runs[0].font.size = Pt(9)
                para.runs[0].bold = i == 0
                if align(sec["table"])[j] == "r":
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for n in sec["notes"]:
            doc.add_paragraph(n, style="List Bullet").runs[0].font.size = Pt(9)
        for p in images(sec):
            doc.add_picture(str(p), width=Inches(5.5))
    doc.add_heading("About this report", 1)
    for a in report["about"]:
        doc.add_paragraph(a, style="List Bullet")
    doc.add_heading("Sources", 1)
    for s in report["sources"]:
        doc.add_paragraph(source_text(s))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ bundles
def _zip(files: dict[str, bytes | str | Path]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in files.items():
            if isinstance(content, Path):
                z.write(content, name)
            else:
                z.writestr(name, content)
    return buf.getvalue()


def _charts(report: dict, folder: str, kinds: set[str] | None = None) -> dict[str, Path]:
    return {f"{folder}/{Path(p).name}": Path(p) for sec in report["sections"] for p in sec.get("charts", [])
            if Path(p).exists() and (kinds is None or Path(p).suffix in kinds)}


def data_zip(report: dict, wb) -> bytes:
    import report_recipes
    import web  # the chat's data download, reused: full series + what was fetched
    record = {"question": report["title"], "asked": report["created"], "evidence": {"results": report["results"]}}
    return web.data_zip(record, "both", wb, report_recipes._loaded.get("longrun"), report_recipes._loaded.get("gdl"))


FORMATS = {  # name -> (file name, media type)
    "pdf": ("report.pdf", "application/pdf"),
    "docx": ("report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "md": ("report_markdown.zip", "application/zip"),
    "tex": ("report_latex.zip", "application/zip"),
    "bib": ("references.bib", "application/x-bibtex"),
    "data": ("data.zip", "application/zip"),
}


def export(report: dict, fmt: str, wb=None) -> bytes:
    if fmt == "pdf":
        return pdf(report)
    if fmt == "docx":
        return docx(report)
    if fmt == "md":
        return _zip({"report.md": markdown(report), "references.bib": bibtex(report), **_charts(report, "charts")})
    if fmt == "tex":
        return _zip({"report.tex": latex(report), "references.bib": bibtex(report),
                     **_charts(report, "figures", {".png"})})
    if fmt == "bib":
        return bibtex(report).encode()
    if fmt == "data":
        return data_zip(report, wb)
    raise ValueError(f"unknown format {fmt!r}; choose from {', '.join(FORMATS)}")


def write_all(report: dict, out_dir: Path, wb=None) -> Path:
    """Every format into reports/<date>_<title>/ (the command-line version)."""
    import workflows
    folder = out_dir / f"{report['created'][:10]}_{workflows.slug(report['title'])}"
    folder.mkdir(parents=True, exist_ok=True)
    for fmt, (name, _) in FORMATS.items():
        if fmt == "data" and wb is None:
            from worldbank import WorldBank
            wb = WorldBank()
        (folder / name).write_bytes(export(report, fmt, wb))
    return folder
